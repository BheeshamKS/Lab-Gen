"""
Document Builder & Styler for LabGenius.
Compiles clean, university-exact Word reports (.docx) matching the student's submission standards:
Title cover page, Program headings, Code screenshots, and Output screenshots.
"""

from pathlib import Path
from typing import List, Optional, Tuple
import base64
import subprocess
import html
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from .config import Config
from .parser import LabManual
from .solver import TaskSolution
from .runner import ExecutionResult


class DocumentBuilder:
    """Builds university-exact lab report documents."""

    def __init__(self, config: Config):
        self.config = config
        self.student = config.student
        self.doc_conf = config.document

    def build_report(
        self,
        manual: LabManual,
        solutions: List[TaskSolution],
        exec_results: List[ExecutionResult],
        task_screenshots: List[Tuple[Path, Path]],  # List of (code_img, out_img)
        output_path: Path,
    ) -> Path:
        """Generate the completed lab document (.docx) matching user's completed lab format."""
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = docx.Document()

        # Set 1.0 inch standard margins
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # -------------------------------------------------------------
        # PAGE 1: OFFICIAL DAWOOD UNIVERSITY COVER PAGE (from first-page.doc)
        # -------------------------------------------------------------
        include_logo = getattr(self.student, "include_logo", True)
        if include_logo:
            custom_logo = getattr(self.student, "custom_logo_path", None)
            if custom_logo and Path(custom_logo).exists():
                logo_path = Path(custom_logo)
            else:
                logo_path = Path(__file__).resolve().parent / "assets" / "university_logo.png"
                if not logo_path.exists():
                    logo_path = Path(__file__).resolve().parent.parent / "web" / "university_logo.png"

            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_logo.paragraph_format.space_before = Pt(36)
            p_logo.paragraph_format.space_after = Pt(28)

            if logo_path and logo_path.exists():
                run_logo = p_logo.add_run()
                run_logo.add_picture(str(logo_path), width=Inches(3.2))

        # Metadata Table matching first-page.doc:
        # Rows: Name, Roll Number, Section, Lab Title, Subject (and Instructor only if toggled on)
        section_val = getattr(self.student, "section", "")
        if not section_val:
            if "-" in self.student.roll_number:
                parts = self.student.roll_number.split("-")
                section_val = f"{parts[0]}-{parts[1]}"
            else:
                section_val = "A"

        subject_val = manual.course_name or self.student.department or "Data Science"
        lab_title_val = self._format_lab_title(manual)

        table_data = [
            ("Name", self.student.name),
            ("Roll Number", self.student.roll_number),
            ("Section", section_val),
            ("Lab Title", lab_title_val),
            ("Subject", subject_val),
        ]

        # Check instructor toggle: only include if explicitly toggled on and provided
        include_inst = getattr(self.student, "include_instructor", False)
        if include_inst and self.student.instructor and self.student.instructor.strip():
            table_data.append(("Instructor", self.student.instructor.strip()))

        cover_table = doc.add_table(rows=len(table_data), cols=2)
        cover_table.style = 'Table Grid'
        cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER

        col_widths = [Inches(2.4), Inches(4.0)]

        for r_idx, (lbl, val) in enumerate(table_data):
            row = cover_table.rows[r_idx]

            c0 = row.cells[0]
            c0.width = col_widths[0]
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_before = Pt(6)
            p0.paragraph_format.space_after = Pt(6)
            r0 = p0.add_run(lbl)
            r0.font.name = self.doc_conf.font_family or "Calibri"
            r0.font.size = Pt(13)

            c1 = row.cells[1]
            c1.width = col_widths[1]
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_before = Pt(6)
            p1.paragraph_format.space_after = Pt(6)
            r1 = p1.add_run(str(val))
            r1.font.name = self.doc_conf.font_family or "Calibri"
            r1.font.size = Pt(13)

        doc.add_page_break()

        # -------------------------------------------------------------
        # PAGES 2+: ONLY THE ASSIGNED TASKS / PROGRAMS
        # -------------------------------------------------------------
        for idx, task in enumerate(manual.tasks):
            sol = next((s for s in solutions if s.task_id == task.task_id), solutions[idx] if idx < len(solutions) else None)
            res = next((r for r in exec_results if r.task_id == task.task_id), exec_results[idx] if idx < len(exec_results) else None)
            code_img, out_img = task_screenshots[idx] if idx < len(task_screenshots) else (None, None)

            # Program / Query Heading
            prefix = "Query" if task.language == "sql" else "Program"
            p_prog = doc.add_paragraph()
            r_prog = p_prog.add_run(f"{prefix} {task.task_id:02d}:  ")
            r_prog.font.bold = True
            r_prog.font.size = Pt(20.0)
            p_prog.paragraph_format.space_before = Pt(0)
            p_prog.paragraph_format.space_after = Pt(4)

            # Subheading: CODE / QUERY:
            p_c_lbl = doc.add_paragraph()
            r_cl = p_c_lbl.add_run("QUERY:" if task.language == "sql" else "CODE:")
            r_cl.font.bold = True
            r_cl.font.size = Pt(18.0)
            p_c_lbl.paragraph_format.space_before = Pt(0)
            p_c_lbl.paragraph_format.space_after = Pt(4)

            # Code Screenshot Image
            if code_img and code_img.exists():
                p_code_img = doc.add_paragraph()
                p_code_img.paragraph_format.space_before = Pt(0)
                p_code_img.paragraph_format.space_after = Pt(6)
                self._add_image_natural(p_code_img, code_img, max_width_in=6.0, max_height_in=5.2)

            # Subheading: OUTPUT:
            p_o_lbl = doc.add_paragraph()
            r_ol = p_o_lbl.add_run("OUTPUT:")
            r_ol.font.bold = True
            r_ol.font.size = Pt(18.0)
            p_o_lbl.paragraph_format.space_before = Pt(4)
            p_o_lbl.paragraph_format.space_after = Pt(4)

            # Output Screenshot Image
            if out_img and out_img.exists():
                p_out_img = doc.add_paragraph()
                p_out_img.paragraph_format.space_before = Pt(0)
                p_out_img.paragraph_format.space_after = Pt(6)
                self._add_image_natural(p_out_img, out_img, max_width_in=6.0, max_height_in=3.8)

            # Page break between programs (except the last one)
            if idx < len(manual.tasks) - 1:
                doc.add_page_break()

        doc.save(output_path)
        return output_path

    def build_pdf(self, docx_path: Path, output_pdf_path: Optional[Path] = None) -> Optional[Path]:
        """Convert completed Word document (.docx) to PDF using headless LibreOffice."""
        if not docx_path.exists():
            return None
        out_dir = output_pdf_path.parent if output_pdf_path else docx_path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        try:
            res = subprocess.run(
                ["libreoffice", "--headless", "--convert-to", "pdf", str(docx_path), "--outdir", str(out_dir)],
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=45,
            )
            expected_pdf = out_dir / f"{docx_path.stem}.pdf"
            if expected_pdf.exists():
                if output_pdf_path and expected_pdf != output_pdf_path:
                    expected_pdf.rename(output_pdf_path)
                    return output_pdf_path
                return expected_pdf
        except Exception as e:
            print(f"[-] PDF conversion error: {e}")
        return None

    def build_markdown_report(
        self,
        manual: LabManual,
        solutions: List[TaskSolution],
        exec_results: List[ExecutionResult],
        task_screenshots: List[Tuple[Path, Path]],
        output_path: Path,
    ) -> Path:
        """Generates a structured, clean Markdown report of the lab submission."""
        output_path.parent.mkdir(parents=True, exist_ok=True)
        lines = []

        section_val = getattr(self.student, "section", "")
        if not section_val:
            if "-" in self.student.roll_number:
                parts = self.student.roll_number.split("-")
                section_val = f"{parts[0]}-{parts[1]}"
            else:
                section_val = "A"

        subject_val = manual.course_name or self.student.department or "Data Science"
        lab_title_val = self._format_lab_title(manual)
        include_inst = getattr(self.student, "include_instructor", False)

        lines.append("# DAWOOD UNIVERSITY OF ENGINEERING & TECHNOLOGY\n")
        lines.append("| Field | Details |")
        lines.append("| :--- | :--- |")
        lines.append(f"| **Name** | {self.student.name} |")
        lines.append(f"| **Roll Number** | {self.student.roll_number} |")
        lines.append(f"| **Section** | {section_val} |")
        lines.append(f"| **Lab Title** | {lab_title_val} |")
        lines.append(f"| **Subject** | {subject_val} |")
        if include_inst and self.student.instructor and self.student.instructor.strip():
            lines.append(f"| **Instructor** | {self.student.instructor.strip()} |")
        lines.append("\n---\n")

        for idx, task in enumerate(manual.tasks):
            sol = next((s for s in solutions if s.task_id == task.task_id), solutions[idx] if idx < len(solutions) else None)
            res = next((r for r in exec_results if r.task_id == task.task_id), exec_results[idx] if idx < len(exec_results) else None)
            code_img, out_img = task_screenshots[idx] if idx < len(task_screenshots) else (None, None)

            prefix = "Query" if task.language == "sql" else "Program"
            lines.append(f"## {prefix} {task.task_id:02d}: {task.title}\n")
            if task.description:
                lines.append(f"> {task.description}\n")

            lines.append(f"### {'QUERY' if task.language == 'sql' else 'CODE'}:\n")
            if sol and sol.code:
                lines.append(f"```{task.language}\n{sol.code}\n```\n")

            if code_img and code_img.exists():
                lines.append(f"![Code Screenshot]({code_img.name})\n")

            lines.append(f"### OUTPUT:\n")
            if res and res.stdout:
                lines.append(f"```text\n{res.stdout.strip()}\n```\n")

            if out_img and out_img.exists():
                lines.append(f"![Output Screenshot]({out_img.name})\n")

            if sol and sol.explanation:
                lines.append(f"**Explanation:** {sol.explanation}\n")

            if sol and sol.discussion_answers:
                lines.append(f"### Discussion / Viva Questions:\n")
                for q, a in sol.discussion_answers.items():
                    lines.append(f"- **Q: {q}**  \n  **A:** {a}\n")

            lines.append("---\n")

        output_path.write_text("\n".join(lines), encoding="utf-8")
        return output_path

    def build_html_report(
        self,
        manual: LabManual,
        solutions: List[TaskSolution],
        exec_results: List[ExecutionResult],
        task_screenshots: List[Tuple[Path, Path]],
        output_path: Path,
        embed_base64: bool = True,
    ) -> Path:
        """Generates a standalone, print-ready HTML report with embedded styles and images."""
        output_path.parent.mkdir(parents=True, exist_ok=True)

        tasks_html = []
        for idx, task in enumerate(manual.tasks):
            sol = next((s for s in solutions if s.task_id == task.task_id), solutions[idx] if idx < len(solutions) else None)
            res = next((r for r in exec_results if r.task_id == task.task_id), exec_results[idx] if idx < len(exec_results) else None)
            code_img, out_img = task_screenshots[idx] if idx < len(task_screenshots) else (None, None)

            code_src = ""
            if code_img and code_img.exists():
                code_src = self._img_to_data_uri(code_img) if embed_base64 else f"/workspace/task_{task.task_id:02d}/{code_img.name}"

            out_src = ""
            if out_img and out_img.exists():
                out_src = self._img_to_data_uri(out_img) if embed_base64 else f"/workspace/task_{task.task_id:02d}/{out_img.name}"

            prefix = "Query" if task.language == "sql" else "Program"
            code_val = html.escape(sol.code) if sol and sol.code else ""
            out_val = html.escape(res.stdout.strip()) if res and res.stdout else ""
            exp_val = html.escape(sol.explanation) if sol and sol.explanation else ""

            viva_html = ""
            if sol and sol.discussion_answers:
                viva_items = "".join(
                    f"<div class='viva-item'><strong>Q: {html.escape(q)}</strong><p>A: {html.escape(a)}</p></div>"
                    for q, a in sol.discussion_answers.items()
                )
                viva_html = f"<div class='viva-section'><h4>Discussion & Viva Questions:</h4>{viva_items}</div>"

            tasks_html.append(f"""
            <section class="report-page">
              <h2 class="prog-heading">{prefix} {task.task_id:02d}: {html.escape(task.title)}</h2>
              {f"<p class='prog-desc'>{html.escape(task.description)}</p>" if task.description else ""}

              <h3 class="section-label">{"QUERY:" if task.language == 'sql' else "CODE:"}</h3>
              {f"<div class='img-container'><img src='{code_src}' alt='Code Screenshot' /></div>" if code_src else ""}
              <div class="code-box"><pre><code>{code_val}</code></pre></div>

              <h3 class="section-label">OUTPUT:</h3>
              {f"<div class='img-container'><img src='{out_src}' alt='Output Screenshot' /></div>" if out_src else ""}
              {f"<div class='console-box'><pre>{out_val}</pre></div>" if out_val else ""}

              {f"<div class='exp-box'><strong>Explanation:</strong> {exp_val}</div>" if exp_val else ""}
              {viva_html}
            </section>
            """)

        include_logo = getattr(self.student, "include_logo", True)
        logo_html = ""
        if include_logo:
            custom_logo = getattr(self.student, "custom_logo_path", None)
            if custom_logo and Path(custom_logo).exists():
                logo_path = Path(custom_logo)
            else:
                logo_path = Path(__file__).resolve().parent / "assets" / "university_logo.png"
                if not logo_path.exists():
                    logo_path = Path(__file__).resolve().parent.parent / "web" / "university_logo.png"

            logo_src = self._img_to_data_uri(logo_path) if (embed_base64 and logo_path and logo_path.exists()) else "/static/university_logo.png"
            logo_html = f'''      <div class="cover-logo-wrapper">
        <img src="{logo_src}" alt="University Logo" class="cover-logo" />
      </div>'''

        section_val = getattr(self.student, "section", "")
        if not section_val:
            if "-" in self.student.roll_number:
                parts = self.student.roll_number.split("-")
                section_val = f"{parts[0]}-{parts[1]}"
            else:
                section_val = "A"

        subject_val = manual.course_name or self.student.department or "Data Science"
        lab_title_val = self._format_lab_title(manual)
        include_inst = getattr(self.student, "include_instructor", False)

        full_html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <title>LAB #{manual.lab_number} — {html.escape(manual.title)}</title>
  <style>
    @page {{
      size: letter;
      margin: 1.0in;
    }}
    body {{
      font-family: 'Calibri', -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif;
      color: #1e293b;
      background: #f8fafc;
      line-height: 1.5;
      margin: 0;
      padding: 24px;
    }}
    .document-wrapper {{
      max-width: 850px;
      margin: 0 auto;
      background: #ffffff;
      padding: 60px;
      box-shadow: 0 10px 25px -5px rgba(0, 0, 0, 0.1), 0 8px 10px -6px rgba(0, 0, 0, 0.1);
      border-radius: 8px;
    }}
    .cover-page {{
      text-align: center;
      padding: 60px 20px 80px;
      page-break-after: always;
      break-after: page;
      border-bottom: 2px dashed #cbd5e1;
    }}
    .cover-logo-wrapper {{
      margin-bottom: 32px;
      text-align: center;
    }}
    .cover-logo {{
      width: 240px;
      height: auto;
    }}
    .cover-table {{
      width: 88%;
      margin: 0 auto;
      border-collapse: collapse;
      font-size: 13pt;
      color: #0f172a;
    }}
    .cover-table td {{
      border: 1px solid #334155;
      padding: 10px 16px;
      text-align: left;
    }}
    .cover-table .lbl-col {{
      width: 35%;
      font-weight: 600;
      background: #f8fafc;
    }}
    .cover-table .val-col {{
      width: 65%;
    }}
    .report-page {{
      page-break-after: always;
      break-after: page;
      padding-top: 30px;
      margin-bottom: 40px;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 40px;
    }}
    .report-page:last-child {{
      border-bottom: none;
      page-break-after: auto;
      break-after: auto;
    }}
    .prog-heading {{
      font-size: 18pt;
      font-weight: 800;
      color: #0f172a;
      margin-bottom: 6px;
    }}
    .prog-desc {{
      color: #64748b;
      font-size: 11pt;
      margin-bottom: 16px;
    }}
    .section-label {{
      font-size: 14pt;
      font-weight: 700;
      color: #1e293b;
      margin-top: 20px;
      margin-bottom: 8px;
    }}
    .img-container {{
      margin: 12px 0 16px;
      text-align: center;
    }}
    .img-container img {{
      max-width: 100%;
      height: auto;
      border: 1px solid #cbd5e1;
      border-radius: 4px;
      box-shadow: 0 2px 6px rgba(0,0,0,0.06);
    }}
    .code-box {{
      background: #f1f5f9;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      padding: 12px 16px;
      font-family: 'Fira Code', 'Consolas', monospace;
      font-size: 10pt;
      overflow-x: auto;
      margin-bottom: 16px;
    }}
    .console-box {{
      background: #0f172a;
      color: #f8fafc;
      border-radius: 6px;
      padding: 12px 16px;
      font-family: 'Fira Code', 'Consolas', monospace;
      font-size: 9.5pt;
      overflow-x: auto;
      margin-bottom: 16px;
    }}
    pre {{
      margin: 0;
      white-space: pre-wrap;
      word-break: break-word;
    }}
    .exp-box {{
      background: #f0fdf4;
      border-left: 4px solid #22c55e;
      padding: 10px 14px;
      border-radius: 4px;
      font-size: 11pt;
      margin-top: 14px;
      color: #166534;
    }}
    .viva-section {{
      margin-top: 20px;
      background: #f8fafc;
      border: 1px solid #e2e8f0;
      border-radius: 6px;
      padding: 14px;
    }}
    .viva-item {{
      margin-bottom: 8px;
    }}
    @media print {{
      body {{
        background: #ffffff;
        padding: 0;
      }}
      .document-wrapper {{
        box-shadow: none;
        padding: 0;
        max-width: 100%;
      }}
      .cover-page {{
        border-bottom: none;
        padding-top: 80px;
      }}
    }}
  </style>
</head>
<body>
  <div class="document-wrapper">
    <div class="cover-page">
{logo_html}

      <table class="cover-table">
        <tbody>
          <tr>
            <td class="lbl-col">Name</td>
            <td class="val-col">{html.escape(self.student.name)}</td>
          </tr>
          <tr>
            <td class="lbl-col">Roll Number</td>
            <td class="val-col">{html.escape(self.student.roll_number)}</td>
          </tr>
          <tr>
            <td class="lbl-col">Section</td>
            <td class="val-col">{html.escape(section_val)}</td>
          </tr>
          <tr>
            <td class="lbl-col">Lab Title</td>
            <td class="val-col">{html.escape(lab_title_val)}</td>
          </tr>
          <tr>
            <td class="lbl-col">Subject</td>
            <td class="val-col">{html.escape(subject_val)}</td>
          </tr>
          {f"<tr><td class='lbl-col'>Instructor</td><td class='val-col'>{html.escape(self.student.instructor.strip())}</td></tr>" if (include_inst and self.student.instructor and self.student.instructor.strip()) else ""}
        </tbody>
      </table>
    </div>

    {''.join(tasks_html)}
  </div>
</body>
</html>
"""
        output_path.write_text(full_html, encoding="utf-8")
        return output_path

    @staticmethod
    def _format_lab_title(manual: LabManual) -> str:
        """Formats the lab title cleanly avoiding duplicate prefixes like Lab #01: Lab #01."""
        if not manual.title:
            return f"Lab #{manual.lab_number}" if manual.lab_number else "Lab Report"
        title_clean = manual.title.strip()
        lower = title_clean.lower()
        if lower.startswith("lab #") or lower.startswith("lab ") or lower.startswith("lab#"):
            return title_clean
        return f"Lab #{manual.lab_number}: {title_clean}" if manual.lab_number else title_clean

    @staticmethod
    def _img_to_data_uri(image_path: Path) -> str:
        """Helper to convert an image file to a base64 data URI for portable HTML viewing."""
        try:
            if not image_path.exists():
                return ""
            ext = image_path.suffix.lstrip(".").lower() or "png"
            with open(image_path, "rb") as f:
                b64 = base64.b64encode(f.read()).decode("ascii")
            return f"data:image/{ext};base64,{b64}"
        except Exception:
            return ""

    @staticmethod
    def _add_image_natural(paragraph, image_path: Path, max_width_in: float = 6.0, max_height_in: float = 5.2):
        """Inserts image preserving 96 DPI natural screen snip scale up to page boundaries."""
        try:
            from PIL import Image
            with Image.open(image_path) as im:
                w_px, h_px = im.size
            
            w_in = w_px / 96.0
            h_in = h_px / 96.0
            
            scale = 1.0
            if w_in > max_width_in:
                scale = min(scale, max_width_in / w_in)
            if h_in > max_height_in:
                scale = min(scale, max_height_in / h_in)
                
            final_w = w_in * scale
            paragraph.add_run().add_picture(str(image_path), width=Inches(final_w))
        except Exception:
            paragraph.add_run().add_picture(str(image_path), width=Inches(min(5.5, max_width_in)))
