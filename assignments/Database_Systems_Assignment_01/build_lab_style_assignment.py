import os
import sys
from pathlib import Path
from PIL import Image

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

# Ensure labgenius is importable
repo_root = Path("/mnt/SSD_Storage/Coding/Lab-Gen")
if str(repo_root) not in sys.path:
    sys.path.insert(0, str(repo_root))

from labgenius.config import Config
from labgenius.sanitizer import DocumentSanitizer

def add_image_natural(paragraph, img_path: Path, max_width_in=6.0, max_height_in=5.0):
    if not img_path.exists():
        print(f"Warning: image {img_path} does not exist.")
        return
    im = Image.open(img_path)
    w_px, h_px = im.size
    dpi = 96.0
    w_in = w_px / dpi
    h_in = h_px / dpi
    scale = min(max_width_in / w_in, max_height_in / h_in, 1.0)
    final_w = Inches(w_in * scale)
    final_h = Inches(h_in * scale)
    run = paragraph.add_run()
    run.add_picture(str(img_path), width=final_w, height=final_h)

def build_lab_style_document():
    config = Config.load(repo_root / "config.yaml")
    doc = docx.Document()

    # 1.0 inch standard margins matching user's labs
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # ------------------------------------------------------------------
    # COVER PAGE: Matches first-page.doc & DBMS_LAB01_completed.docx
    # ------------------------------------------------------------------
    logo_path = repo_root / "labgenius" / "assets" / "university_logo.png"
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(36)
    p_logo.paragraph_format.space_after = Pt(28)

    if logo_path.exists():
        r_logo = p_logo.add_run()
        r_logo.add_picture(str(logo_path), width=Inches(3.2))

    # Cover table matching first-page.doc
    table_data = [
        ("Name", "Bheesham Kumar Sajnani"),
        ("Roll Number", "25F-DS-020"),
        ("Section", "25F-DS"),
        ("Lab Title", "Database Systems — Assignment 01"),
        ("Subject", "Database Systems"),
        ("Instructor", "Engr. Seema Bughio")
    ]

    cover_table = doc.add_table(rows=len(table_data), cols=2)
    cover_table.style = 'Table Grid'
    cover_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    col_widths = [Inches(2.4), Inches(4.0)]
    for r_idx, (lbl, val) in enumerate(table_data):
        row = cover_table.rows[r_idx]
        c0 = row.cells[0]
        c0.width = col_widths[0]
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_before = Pt(6)
        p0.paragraph_format.space_after = Pt(6)
        r0 = p0.add_run(lbl)
        r0.font.name = "Calibri"
        r0.font.size = Pt(13)
        r0.bold = True

        c1 = row.cells[1]
        c1.width = col_widths[1]
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_before = Pt(6)
        p1.paragraph_format.space_after = Pt(6)
        r1 = p1.add_run(val)
        r1.font.name = "Calibri"
        r1.font.size = Pt(13)

    # Page break after cover page
    doc.add_page_break()

    # Base image directory
    img_dir = repo_root / "assignments" / "Database_Systems_Assignment_01" / "images"

    # ------------------------------------------------------------------
    # TASK 01: Database Role in Information Systems
    # ------------------------------------------------------------------
    p_t1 = doc.add_paragraph()
    r_t1 = p_t1.add_run("Task 01: Database Role in Information Systems")
    r_t1.bold = True
    r_t1.font.size = Pt(14.0)
    p_t1.paragraph_format.space_before = Pt(4)
    p_t1.paragraph_format.space_after = Pt(4)

    p_d1 = doc.add_paragraph()
    p_d1.paragraph_format.space_before = Pt(0)
    p_d1.paragraph_format.space_after = Pt(6)
    r_dl1 = p_d1.add_run("Requirements / Details: ")
    r_dl1.bold = True
    r_dl1.font.size = Pt(11.0)
    r_dt1 = p_d1.add_run("Explain how a database supports a University Information System. Identify at least four types of information that can be stored in the database and explain how different users can use this information.")
    r_dt1.font.size = Pt(10.5)

    p_ans1 = doc.add_paragraph()
    r_ans1 = p_ans1.add_run("Solution:")
    r_ans1.bold = True
    r_ans1.font.size = Pt(11.5)
    p_ans1.paragraph_format.space_before = Pt(2)
    p_ans1.paragraph_format.space_after = Pt(4)

    p_body1 = doc.add_paragraph(
        "A University Information System (UIS) relies on a centralized Relational Database Management System (DBMS) to operate reliably. "
        "Instead of keeping separate files or Excel sheets across departments which cause data mismatch and duplication, the database acts as a single central source of truth. "
        "It enforces referential integrity constraints so invalid data cannot be entered, allows hundreds of students and faculty to access the system at the same time without data loss, and restricts confidential data (like exam papers and salaries) using user permissions."
    )
    p_body1.paragraph_format.space_after = Pt(6)

    # Embed Architecture Diagram
    uis_diag = img_dir / "uis_architecture.png"
    if uis_diag.exists():
        p_diag = doc.add_paragraph()
        p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_diag.paragraph_format.space_before = Pt(4)
        p_diag.paragraph_format.space_after = Pt(8)
        add_image_natural(p_diag, uis_diag, max_width_in=6.0, max_height_in=3.6)

    # 4 Types of Information table
    p_info_title = doc.add_paragraph()
    r_it = p_info_title.add_run("Types of Information Stored in University Database:")
    r_it.bold = True
    r_it.font.size = Pt(11.0)

    info_table = doc.add_table(rows=5, cols=3)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_headers = ["#", "Information Type", "Attributes Stored"]
    for c_i, h_text in enumerate(info_headers):
        cell = info_table.cell(0, c_i)
        p = cell.paragraphs[0]
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(10)
    
    info_rows = [
        ("1", "Student Records", "Student ID, Name, Department, Enrollment Year, Courses Enrolled, Attendance, CGPA."),
        ("2", "Faculty Records", "Faculty ID, Name, Designation, Department, Assigned Courses, Timetable, Salary Info."),
        ("3", "Course Catalog", "Course Code (e.g. DS-211), Title, Credit Hours, Prerequisites, Syllabus, Allocated Lab."),
        ("4", "Fee Transactions", "Voucher Number, Student ID, Due Date, Tuition Amount, Payment Status (Paid/Pending), Bank ID.")
    ]
    for r_i, (n, cat, desc) in enumerate(info_rows, start=1):
        for c_i, val in enumerate([n, cat, desc]):
            cell = info_table.cell(r_i, c_i)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if c_i == 1:
                r.bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    p_users_title = doc.add_paragraph()
    r_ut = p_users_title.add_run("How Different Users Use This Information:")
    r_ut.bold = True
    r_ut.font.size = Pt(11.0)

    user_points = [
        ("Students: ", "Check timetable, register for semester courses, view attendance, check exam grades and CGPA, and download fee vouchers."),
        ("Faculty / Lecturers: ", "Access enrolled class rosters, enter student lecture attendance, upload quiz/assignment/exam marks, and compute internal assessments."),
        ("Examination Cell: ", "Verify student grades, calculate grade distribution curves, verify prerequisite clearance, and print official transcripts."),
        ("Accounts Department: ", "Verify fee payments from banks, generate student clearance for admit cards, and process staff payroll.")
    ]
    for u_title, u_desc in user_points:
        bp = doc.add_paragraph(style='List Bullet')
        bp.paragraph_format.space_before = Pt(1)
        bp.paragraph_format.space_after = Pt(2)
        r_bt = bp.add_run(u_title)
        r_bt.bold = True
        bp.add_run(u_desc)

    # ------------------------------------------------------------------
    # TASK 02: Basic Database Security
    # ------------------------------------------------------------------
    p_t2 = doc.add_paragraph()
    r_t2 = p_t2.add_run("Task 02: Basic Database Security")
    r_t2.bold = True
    r_t2.font.size = Pt(14.0)
    p_t2.paragraph_format.space_before = Pt(12)
    p_t2.paragraph_format.space_after = Pt(4)

    p_d2 = doc.add_paragraph()
    p_d2.paragraph_format.space_before = Pt(0)
    p_d2.paragraph_format.space_after = Pt(6)
    r_dl2 = p_d2.add_run("Requirements / Details: ")
    r_dl2.bold = True
    r_dl2.font.size = Pt(11.0)
    r_dt2 = p_d2.add_run("A university database contains students' personal information and examination results. Identify three basic security measures that should be implemented to protect this data. Explain the purpose of each measure.")
    r_dt2.font.size = Pt(10.5)

    p_ans2 = doc.add_paragraph()
    r_ans2 = p_ans2.add_run("Solution:")
    r_ans2.bold = True
    r_ans2.font.size = Pt(11.5)
    p_ans2.paragraph_format.space_before = Pt(2)
    p_ans2.paragraph_format.space_after = Pt(4)

    sec_measures = [
        (
            "1. Authentication and Role-Based Access Control (RBAC):",
            "Users must log in with secure credentials (passwords with salted hashing and MFA for staff). The DBMS enforces least privilege: students only get SELECT permission on their own grades via views; teachers get UPDATE/INSERT rights only on courses they teach; admin/exam staff have audited access.",
            "Purpose: ",
            "Prevents unauthorized users or students from accessing sensitive records, modifying marks, or viewing peers' private personal data."
        ),
        (
            "2. Data Encryption (In-Transit and At-Rest):",
            "Network traffic between web portals and the database is secured using TLS 1.3 / HTTPS. At the storage layer, Transparent Data Encryption (TDE with AES-256) encrypts database files and backups.",
            "Purpose: ",
            "Protects data confidentiality. Even if network packets are sniffed or a physical hard drive/backup is stolen, the data cannot be read without cryptographic keys."
        ),
        (
            "3. Database Auditing, Activity Logging & Check Constraints:",
            "An automated audit log records every grade update (User_ID, old marks, new marks, timestamp, and IP). Check constraints (e.g. 0 <= Marks <= 100) ensure only valid values are accepted.",
            "Purpose: ",
            "Guarantees accountability, non-repudiation, and traceability. If an unauthorized grade alteration is attempted, administrators can see exactly who did it and rollback the changes."
        )
    ]

    for m_head, m_body, p_lbl, p_text in sec_measures:
        p_m = doc.add_paragraph()
        p_m.paragraph_format.space_before = Pt(4)
        p_m.paragraph_format.space_after = Pt(2)
        r_mh = p_m.add_run(m_head + " ")
        r_mh.bold = True
        p_m.add_run(m_body)

        p_purp = doc.add_paragraph()
        p_purp.paragraph_format.space_before = Pt(1)
        p_purp.paragraph_format.space_after = Pt(4)
        r_pl = p_purp.add_run(p_lbl)
        r_pl.bold = True
        r_pl.font.color.rgb = RGBColor(180, 83, 9)
        p_purp.add_run(p_text)

    # Embed Security Architecture Diagram
    sec_diag = img_dir / "database_security_layers.png"
    if sec_diag.exists():
        p_sdiag = doc.add_paragraph()
        p_sdiag.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sdiag.paragraph_format.space_before = Pt(4)
        p_sdiag.paragraph_format.space_after = Pt(8)
        add_image_natural(p_sdiag, sec_diag, max_width_in=6.0, max_height_in=3.4)

    # ------------------------------------------------------------------
    # TASK 03: DBMS Operations (SQL Queries)
    # ------------------------------------------------------------------
    p_t3 = doc.add_paragraph()
    r_t3 = p_t3.add_run("Task 03: DBMS Operations (SQL Queries)")
    r_t3.bold = True
    r_t3.font.size = Pt(14.0)
    p_t3.paragraph_format.space_before = Pt(12)
    p_t3.paragraph_format.space_after = Pt(4)

    p_d3 = doc.add_paragraph()
    p_d3.paragraph_format.space_before = Pt(0)
    p_d3.paragraph_format.space_after = Pt(6)
    r_dl3 = p_d3.add_run("Requirements / Details: ")
    r_dl3.bold = True
    r_dl3.font.size = Pt(11.0)
    r_dt3 = p_d3.add_run("Given a Student table containing Student_ID, Name, Department, and CGPA, write SQL queries to perform the following operations: (1) Display students with CGPA > 3.0, (2) Change department of a particular student, (3) Remove a student from the table, (4) Add a new student.")
    r_dt3.font.size = Pt(10.5)

    # Table Setup Sub-section
    p_setup_title = doc.add_paragraph()
    r_st = p_setup_title.add_run("Table Schema & Initial Records Setup:")
    r_st.bold = True
    r_st.font.size = Pt(11.5)

    # Setup Query & Output
    p_c_lbl = doc.add_paragraph()
    r_cl = p_c_lbl.add_run("QUERY:")
    r_cl.font.bold = True
    r_cl.font.size = Pt(13.0)
    p_c_lbl.paragraph_format.space_before = Pt(4)
    p_c_lbl.paragraph_format.space_after = Pt(4)

    p_setup_code_img = doc.add_paragraph()
    p_setup_code_img.paragraph_format.space_before = Pt(0)
    p_setup_code_img.paragraph_format.space_after = Pt(4)
    add_image_natural(p_setup_code_img, img_dir / "q3_setup_code.png", max_width_in=6.0, max_height_in=4.5)

    p_o_lbl = doc.add_paragraph()
    r_ol = p_o_lbl.add_run("OUTPUT:")
    r_ol.font.bold = True
    r_ol.font.size = Pt(13.0)
    p_o_lbl.paragraph_format.space_before = Pt(4)
    p_o_lbl.paragraph_format.space_after = Pt(4)

    p_setup_out_img = doc.add_paragraph()
    p_setup_out_img.paragraph_format.space_before = Pt(0)
    p_setup_out_img.paragraph_format.space_after = Pt(8)
    add_image_natural(p_setup_out_img, img_dir / "q3_setup_output.png", max_width_in=6.0, max_height_in=4.5)

    # 4 Operations in exact lab format
    ops = [
        (
            "Operation 01: Display students with CGPA greater than 3.0",
            "SELECT statement with WHERE clause filtering students where CGPA strictly exceeds 3.0.",
            "q3_op1_code.png",
            "q3_op1_output.png"
        ),
        (
            "Operation 02: Change the department of a particular student",
            "UPDATE statement modifying the Department column for student 'CS-25-045' uniquely by primary key.",
            "q3_op2_code.png",
            "q3_op2_output.png"
        ),
        (
            "Operation 03: Remove a student from the table",
            "DELETE FROM statement removing student 'DS-25-099' uniquely by primary key.",
            "q3_op3_code.png",
            "q3_op3_output.png"
        ),
        (
            "Operation 04: Add a new student",
            "INSERT INTO statement adding a new student record 'DS-25-102' (Ayesha Khan) into the Student table.",
            "q3_op4_code.png",
            "q3_op4_output.png"
        )
    ]

    for op_title, op_desc, code_img_name, out_img_name in ops:
        p_op = doc.add_paragraph()
        r_opt = p_op.add_run(op_title)
        r_opt.bold = True
        r_opt.font.size = Pt(12.0)
        p_op.paragraph_format.space_before = Pt(10)
        p_op.paragraph_format.space_after = Pt(2)

        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(4)
        r_dl = p_desc.add_run("Purpose: ")
        r_dl.bold = True
        r_dl.font.size = Pt(10.5)
        r_dt = p_desc.add_run(op_desc)
        r_dt.font.size = Pt(10.0)

        # QUERY:
        p_c = doc.add_paragraph()
        r_c = p_c.add_run("QUERY:")
        r_c.font.bold = True
        r_c.font.size = Pt(13.0)
        p_c.paragraph_format.space_before = Pt(3)
        p_c.paragraph_format.space_after = Pt(3)

        p_c_img = doc.add_paragraph()
        p_c_img.paragraph_format.space_before = Pt(0)
        p_c_img.paragraph_format.space_after = Pt(4)
        add_image_natural(p_c_img, img_dir / code_img_name, max_width_in=6.0, max_height_in=4.2)

        # OUTPUT:
        p_o = doc.add_paragraph()
        r_o = p_o.add_run("OUTPUT:")
        r_o.font.bold = True
        r_o.font.size = Pt(13.0)
        p_o.paragraph_format.space_before = Pt(3)
        p_o.paragraph_format.space_after = Pt(3)

        p_o_img = doc.add_paragraph()
        p_o_img.paragraph_format.space_before = Pt(0)
        p_o_img.paragraph_format.space_after = Pt(6)
        add_image_natural(p_o_img, img_dir / out_img_name, max_width_in=6.0, max_height_in=4.2)

    # ------------------------------------------------------------------
    # TASK 04: Security Scenario
    # ------------------------------------------------------------------
    p_t4 = doc.add_paragraph()
    r_t4 = p_t4.add_run("Task 04: Security Scenario")
    r_t4.bold = True
    r_t4.font.size = Pt(14.0)
    p_t4.paragraph_format.space_before = Pt(12)
    p_t4.paragraph_format.space_after = Pt(4)

    p_d4 = doc.add_paragraph()
    p_d4.paragraph_format.space_before = Pt(0)
    p_d4.paragraph_format.space_after = Pt(6)
    r_dl4 = p_d4.add_run("Requirements / Details: ")
    r_dl4.bold = True
    r_dl4.font.size = Pt(11.0)
    r_dt4 = p_d4.add_run("A student is able to access and modify another student's examination marks through the university database system. (a) What security problem exists? (b) What database security mechanism could help prevent this problem?")
    r_dt4.font.size = Pt(10.5)

    p_ans4 = doc.add_paragraph()
    r_ans4 = p_ans4.add_run("Solution:")
    r_ans4.bold = True
    r_ans4.font.size = Pt(11.5)
    p_ans4.paragraph_format.space_before = Pt(2)
    p_ans4.paragraph_format.space_after = Pt(4)

    # Part a
    p_4a = doc.add_paragraph()
    r_4ah = p_4a.add_run("Part (a): Security Problem Identified\n")
    r_4ah.bold = True
    r_4ah.font.size = Pt(11.0)
    p_4a.add_run(
        "1. Broken Access Control & Insecure Direct Object Reference (IDOR): The system fails to check if the user requesting the record owns or has permission for it. A student is able to target another student's record and modify it.\n"
        "2. Violation of Principle of Least Privilege: Regular student user accounts have UPDATE and INSERT write permissions on examination marks tables, which should only belong to faculty and the exam cell.\n"
        "3. Breach of Confidentiality and Integrity: Confidentiality is violated because private marks are visible to peers; Integrity is destroyed because unauthorized users can falsify academic grades."
    )
    p_4a.paragraph_format.space_after = Pt(6)

    # Part b
    p_4b = doc.add_paragraph()
    r_4bh = p_4b.add_run("Part (b): Database Security Mechanism to Prevent This\n")
    r_4bh.bold = True
    r_4bh.font.size = Pt(11.0)
    p_4b.add_run(
        "1. DCL Privilege Revocation: Revoke all UPDATE, INSERT, and DELETE permissions on Examination_Marks from the Student_Role.\n"
        "2. Row-Level Security (RLS): Enable RLS policies in the DBMS so students can strictly SELECT rows where Student_ID = CURRENT_USER(). Any query attempting to access or modify another student's row is blocked at the database kernel level.\n"
        "3. Secure Views and Stored Procedures: Force students to query a secure view (View_Student_Own_Marks) instead of the base table, and require teachers to submit grades via validated stored procedures.\n"
        "4. Audit Triggers: Create an AFTER UPDATE trigger on examination marks that automatically logs any changes to an immutable audit table with timestamps."
    )
    p_4b.paragraph_format.space_after = Pt(6)

    # Security Query & Output
    p_sq_lbl = doc.add_paragraph()
    r_sql = p_sq_lbl.add_run("QUERY (Security Implementation):")
    r_sql.font.bold = True
    r_sql.font.size = Pt(13.0)
    p_sq_lbl.paragraph_format.space_before = Pt(3)
    p_sq_lbl.paragraph_format.space_after = Pt(3)

    p_sq_img = doc.add_paragraph()
    p_sq_img.paragraph_format.space_before = Pt(0)
    p_sq_img.paragraph_format.space_after = Pt(4)
    add_image_natural(p_sq_img, img_dir / "q4_sec_code.png", max_width_in=6.0, max_height_in=4.2)

    p_so_lbl = doc.add_paragraph()
    r_sol = p_so_lbl.add_run("OUTPUT:")
    r_sol.font.bold = True
    r_sol.font.size = Pt(13.0)
    p_so_lbl.paragraph_format.space_before = Pt(3)
    p_so_lbl.paragraph_format.space_after = Pt(3)

    p_so_img = doc.add_paragraph()
    p_so_img.paragraph_format.space_before = Pt(0)
    p_so_img.paragraph_format.space_after = Pt(6)
    add_image_natural(p_so_img, img_dir / "q4_sec_output.png", max_width_in=6.0, max_height_in=4.2)

    # Save DOCX
    docx_out = repo_root / "assignments" / "Database_Systems_Assignment_01" / "Database_Systems_Assignment_01.docx"
    doc.save(str(docx_out))
    print(f"DOCX saved to: {docx_out}")

    # Sanitize metadata using LabGenius sanitizer
    sanitizer = DocumentSanitizer(config)
    sanitizer.sanitize_docx(docx_out)
    print("DOCX metadata sanitized with human student profile (Bheesham Kumar Sajnani)!")

    # Convert to PDF using headless LibreOffice
    pdf_out = docx_out.with_suffix(".pdf")
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(docx_out), "--outdir", str(docx_out.parent)]
    subprocess.run(cmd, check=True)
    print(f"PDF generated successfully at: {pdf_out}")

if __name__ == "__main__":
    import subprocess
    build_lab_style_document()
