"""
Generator for Held-Out Evaluation Corpus.
Creates genuinely unseen lab manuals in DOCX, PDF, TXT, and Markdown formats,
and produces evaluation_corpus/manifest.json separating Development, Regression,
and Held-Out fixtures.
"""

import os
import json
import subprocess
from pathlib import Path
import docx
from docx.shared import Pt, Inches

CORPUS_DIR = Path(__file__).resolve().parent
HELD_OUT_DIR = CORPUS_DIR / "held_out"
HELD_OUT_DIR.mkdir(parents=True, exist_ok=True)


def build_manual_1_docx():
    """Manual 1: DSA - Doubly Linked List (DOCX format)"""
    file_path = HELD_OUT_DIR / "held_out_01_dsa_doubly_linked_list.docx"
    doc = docx.Document()

    doc.add_heading("Department of Computer Science & Software Engineering", level=1)
    doc.add_heading("Course: Data Structures and Algorithms", level=2)
    doc.add_heading("Lab Manual - Session 07", level=2)
    p_title = doc.add_paragraph("Topic: Doubly Linked List Operations & Bidirectional Navigation")
    p_title.runs[0].font.bold = True

    doc.add_heading("1. Objectives", level=3)
    doc.add_paragraph("1. Understand the memory layout and pointer architecture of doubly linked lists.")
    doc.add_paragraph("2. Implement node insertion at head and tail.")
    doc.add_paragraph("3. Implement node deletion by target value with boundary check.")
    doc.add_paragraph("4. Perform forward and backward traversals.")

    doc.add_heading("2. Theoretical Background", level=3)
    doc.add_paragraph(
        "A Doubly Linked List is a linear data structure where each node consists of three fields: "
        "data, next pointer, and prev pointer. This enables bidirectional navigation at the expense "
        "of increased pointer manipulation overhead."
    )

    doc.add_heading("3. LAB EXERCISES:", level=3)
    doc.add_paragraph("1. Insertion Operations: Implement a DoublyLinkedList class with insert_at_head and insert_at_tail methods in Python.")
    doc.add_paragraph("2. Node Deletion Mechanism: Implement a delete_node method that searches for a value and removes it, updating neighboring pointers correctly.")
    doc.add_paragraph("3. Bidirectional List Traversal: Implement display_forward and display_backward methods to traverse and print the list in both directions.")

    doc.add_heading("4. Questions for Analysis:", level=3)
    doc.add_paragraph("1. Analyze the time and space complexity of Doubly Linked List insertion compared to Singly Linked List.")
    doc.add_paragraph("2. Detail the exact pointer edge cases encountered when deleting the only remaining node in a Doubly Linked List.")

    doc.add_heading("5. Viva Voce Questions:", level=3)
    doc.add_paragraph("1. What is the memory overhead of maintaining the prev pointer in a doubly linked list?")
    doc.add_paragraph("2. Under what practical operating system or application conditions is a doubly linked list preferred over an array?")

    doc.save(file_path)
    print(f"Created: {file_path}")
    return file_path


def build_manual_2_pdf():
    """Manual 2: DBMS - Library Relational Schema (PDF format via LibreOffice)"""
    docx_temp = HELD_OUT_DIR / "_temp_held_out_02.docx"
    pdf_path = HELD_OUT_DIR / "held_out_02_dbms_library_system.pdf"
    doc = docx.Document()

    doc.add_heading("Faculty of Computing & Information Technology", level=1)
    doc.add_heading("Course: Database Management Systems (DBMS)", level=2)
    doc.add_heading("Laboratory Manual - Experiment 08", level=2)
    doc.add_paragraph("Title: Relational Library Schema & Circulation Tracking")

    doc.add_heading("1. Relational Schema Specification", level=3)
    doc.add_paragraph("Books (BookID INT PRIMARY KEY, Title VARCHAR(100), Author VARCHAR(100), Category VARCHAR(50), Price DECIMAL(10,2))")
    doc.add_paragraph("Members (MemberID INT PRIMARY KEY, FullName VARCHAR(100), Email VARCHAR(100), JoinDate DATE)")
    doc.add_paragraph("Loans (LoanID INT PRIMARY KEY, BookID INT REFERENCES Books(BookID), MemberID INT REFERENCES Members(MemberID), LoanDate DATE, ReturnDate DATE)")

    doc.add_heading("2. LAB TASKS:", level=3)
    doc.add_paragraph("Query 1: Relational Table Creation with Constraints: Write SQL statements to create Books, Members, and Loans tables with PRIMARY KEY and FOREIGN KEY constraints.")
    doc.add_paragraph("Query 2: Book Circulation Category Aggregation: Write a query grouping books by Category, counting total books and computing average price, displaying only categories with at least 1 book.")
    doc.add_paragraph("Query 3: Member Active Borrowing Join Report: Write an INNER JOIN query between Members and Loans to retrieve member names, book IDs, and loan dates for active borrowings.")

    doc.add_heading("3. Analysis Questions:", level=3)
    doc.add_paragraph("1. How do foreign key constraints enforce referential integrity across the Books and Loans tables?")
    doc.add_paragraph("2. Compare the query execution plan of an INNER JOIN versus subquery with IN clause.")

    doc.add_heading("4. Viva Voce Questions:", level=3)
    doc.add_paragraph("1. What is the fundamental difference between WHERE and HAVING clauses in SQL?")
    doc.add_paragraph("2. Can a table have multiple foreign keys referencing different tables?")

    doc.save(docx_temp)

    # Convert to PDF with LibreOffice
    cmd = ["libreoffice", "--headless", "--convert-to", "pdf", str(docx_temp), "--outdir", str(HELD_OUT_DIR)]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=30)
    converted_pdf = HELD_OUT_DIR / f"{docx_temp.stem}.pdf"
    if converted_pdf.exists():
        converted_pdf.rename(pdf_path)
    if docx_temp.exists():
        docx_temp.unlink()

    print(f"Created: {pdf_path}")
    return pdf_path


def build_manual_3_txt():
    """Manual 3: Systems C - String Algorithms (TXT format)"""
    file_path = HELD_OUT_DIR / "held_out_03_systems_c_strings.txt"
    content = """Session: 09 - Low-Level Systems Programming Lab
Course: Systems Programming in C
Title: C STRING ALGORITHMS WITHOUT STANDARD LIBRARY

OBJECTIVES:
1. Implement pointer-based character array manipulation.
2. Prevent standard string library dependencies (string.h).
3. Handle null-terminated strings and buffer bounds.

ASSIGNMENT TASKS:
Task 1: Palindrome Verification
Write a C program that checks whether an input string is a palindrome using two-pointer approach without using string.h functions.

Task 2: Sentence Word Counter
Write a C program to count total words in a sentence by scanning whitespace delimiters and consecutive non-space tokens.

Task 3: Substring Search Engine
Write a C program to search for a target substring in a source string and return the starting 0-based index or -1 if not found.

QUESTIONS FOR ANALYSIS:
1. Discuss the security vulnerabilities of gets() versus fgets() when reading character arrays in C.
2. What is the algorithmic time complexity of naive substring search versus Knuth-Morris-Pratt?

VIVA QUESTIONS:
1. Why must all C strings end with the null terminator character '\\0'?
2. What happens if a C program attempts to access an index beyond the allocated array size?
"""
    file_path.write_text(content, encoding="utf-8")
    print(f"Created: {file_path}")
    return file_path


def build_manual_4_md():
    """Manual 4: Data Science - Pandas Outlier Detection (Markdown format)"""
    file_path = HELD_OUT_DIR / "held_out_04_data_science_pandas.md"
    content = """# Data Science Laboratory Manual
## Session: 10 - Exploratory Data Analysis and Statistical Summaries
**Course:** Data Science & Applied Statistics  
**Topic:** Transaction Data Distribution Analysis and IQR Outlier Detection

### Scenario & Dataset Description:
The dataset simulates electronic banking transactions containing transaction identifier, customer age, transaction amount ($), and customer rating.

### LAB TASKS:
#### Task 1: Descriptive Statistics Summary
Write a Python script using pandas to calculate the mean, median, standard deviation, and variance for transaction amounts from sample transaction records.

#### Task 2: Outlier Detection using Interquartile Range
Write a Python script to compute the 25th percentile (Q1), 75th percentile (Q3), and IQR for transaction amounts, identifying all records where amount > Q3 + 1.5 * IQR.

#### Task 3: Summary Visualization and Distribution Reporting
Write a Python script to summarize transaction distribution statistics and display formatted metrics.

### OBSERVATION TABLE:
| Metric | Formula | Expected Calculation |
| :--- | :--- | :--- |
| Mean | sum(x) / N | Arithmetic Average of Amounts |
| Median | Q2 (50th percentile) | Middle Value |
| IQR | Q3 - Q1 | Spread of Middle 50% |

### Questions for Analysis:
1. Under what statistical distribution conditions is the median preferred over the mean as a measure of central tendency?
2. How does the 1.5 * IQR multiplier threshold minimize false positives in financial anomaly detection?

### Viva Voce Questions:
1. What is the definition of the Interquartile Range (IQR)?
2. Why is NumPy or Pandas vectorized operation faster than a standard Python for-loop?
"""
    file_path.write_text(content, encoding="utf-8")
    print(f"Created: {file_path}")
    return file_path


def build_manual_5_malformed_txt():
    """Manual 5: Ambiguous / Malformed Metadata Fixture (TXT format)"""
    file_path = HELD_OUT_DIR / "held_out_05_malformed_metadata.txt"
    content = """ALGORITHMS LAB PRACTICE SHEET
Department of Software Studies

PRACTICE PROBLEMS:
Problem 1: Implement Quick Sort partitioning on an array of integers.
Problem 2: Implement Linear Search to find target element in an unsorted list.

VIVA QUESTIONS:
1. What is the average time complexity of Quick Sort?
2. What is the worst-case complexity of Linear Search?
"""
    file_path.write_text(content, encoding="utf-8")
    print(f"Created: {file_path}")
    return file_path


def build_manifest():
    """Builds evaluation_corpus/manifest.json separating categories."""
    manifest = {
        "metadata": {
            "version": "1.0",
            "description": "Evaluation Corpus Manifest distinguishing Development, Regression, and Genuinely Held-Out manuals.",
        },
        "development_examples": [
            {
                "id": "DEV_01",
                "file": "examples/sample_lab_manual.docx",
                "domain": "DSA",
                "format": "docx",
                "purpose": "Original sample lab manual used for UI and base feature demonstrations.",
            },
            {
                "id": "DEV_02",
                "file": "examples/Lab 03 manual.docx",
                "domain": "DSA",
                "format": "docx",
                "purpose": "Reference Lab 03 manual used for format matching.",
            },
            {
                "id": "DEV_03",
                "file": "uploads/Data Structure Laboratory Manual L1.docx",
                "domain": "DSA",
                "format": "docx",
                "purpose": "Dawood University DSA Lab 01 development reference.",
            },
            {
                "id": "DEV_04",
                "file": "uploads/DBMS Lab 01.docx",
                "domain": "DBMS",
                "format": "docx",
                "purpose": "Dawood University DBMS Lab 01 development reference.",
            },
        ],
        "regression_fixtures": [
            {
                "id": "REG_01",
                "test_file": "tests/test_audit_and_remediation.py",
                "purpose": "First-pass audit fixes regression test suite (50 assertions).",
            },
            {
                "id": "REG_02",
                "test_file": "tests/test_second_pass_adversarial.py",
                "purpose": "Second-pass adversarial security, path traversal, and document inspection tests.",
            },
            {
                "id": "REG_03",
                "test_file": "tests/test_held_out_generalization.py",
                "purpose": "Leave-one-out and synthetic DSA/DBMS generalization regression tests.",
            },
        ],
        "held_out_evaluation_examples": [
            {
                "id": "HELDOUT_01",
                "file": "evaluation_corpus/held_out/held_out_01_dsa_doubly_linked_list.docx",
                "format": "docx",
                "domain": "Data Structures & Algorithms",
                "language": "python",
                "expected_lab_number": "07",
                "expected_tasks_count": 3,
                "expected_tasks": [
                    "Task 1: Insertion Operations",
                    "Task 2: Node Deletion Mechanism",
                    "Task 3: Bidirectional List Traversal",
                ],
                "expected_analysis_count": 2,
                "expected_viva_count": 2,
            },
            {
                "id": "HELDOUT_02",
                "file": "evaluation_corpus/held_out/held_out_02_dbms_library_system.pdf",
                "format": "pdf",
                "domain": "Database Management Systems",
                "language": "sql",
                "expected_lab_number": "08",
                "expected_tasks_count": 3,
                "expected_tasks": [
                    "Query 1: Relational Table Creation with Constraints",
                    "Query 2: Book Circulation Category Aggregation",
                    "Query 3: Member Active Borrowing Join Report",
                ],
                "expected_analysis_count": 2,
                "expected_viva_count": 2,
            },
            {
                "id": "HELDOUT_03",
                "file": "evaluation_corpus/held_out/held_out_03_systems_c_strings.txt",
                "format": "txt",
                "domain": "Systems Programming",
                "language": "c",
                "expected_lab_number": "09",
                "expected_tasks_count": 3,
                "expected_tasks": [
                    "Task 1: Palindrome Verification",
                    "Task 2: Sentence Word Counter",
                    "Task 3: Substring Search Engine",
                ],
                "expected_analysis_count": 2,
                "expected_viva_count": 2,
            },
            {
                "id": "HELDOUT_04",
                "file": "evaluation_corpus/held_out/held_out_04_data_science_pandas.md",
                "format": "md",
                "domain": "Data Science & Statistics",
                "language": "python",
                "expected_lab_number": "10",
                "expected_tasks_count": 3,
                "expected_tasks": [
                    "Task 1: Descriptive Statistics Summary",
                    "Task 2: Outlier Detection using Interquartile Range",
                    "Task 3: Summary Visualization and Distribution Reporting",
                ],
                "expected_analysis_count": 2,
                "expected_viva_count": 2,
                "has_observation_table": True,
            },
            {
                "id": "HELDOUT_05",
                "file": "evaluation_corpus/held_out/held_out_05_malformed_metadata.txt",
                "format": "txt",
                "domain": "General Algorithms",
                "language": "python",
                "expected_lab_number": "01",
                "expected_tasks_count": 2,
                "expected_tasks": [
                    "Problem 1: Quick Sort",
                    "Problem 2: Linear Search",
                ],
                "expected_analysis_count": 0,
                "expected_viva_count": 2,
                "notes": "Malformed / ambiguous metadata fixture without explicit lab number or Title prefix.",
            },
        ],
    }

    manifest_path = CORPUS_DIR / "manifest.json"
    with open(manifest_path, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    print(f"Created: {manifest_path}")


if __name__ == "__main__":
    build_manual_1_docx()
    build_manual_2_pdf()
    build_manual_3_txt()
    build_manual_4_md()
    build_manual_5_malformed_txt()
    build_manifest()
