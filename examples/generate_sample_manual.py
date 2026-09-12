"""
Generates a realistic university lab manual (.docx) for testing and demonstration.
"""

from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def create_sample_manual(dest_path: Path):
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = title_p.add_run("DEPARTMENT OF DATA SCIENCE & COMPUTING\n")
    r_uni.font.bold = True
    r_uni.font.size = Pt(14)
    r_uni.font.color.rgb = RGBColor(33, 73, 114)

    r_course = title_p.add_run("Course: Data Science Tools & Techniques (DS-204)\n")
    r_course.font.bold = True
    r_course.font.size = Pt(12)

    r_title = title_p.add_run("LAB MANUAL 03: STATISTICAL MODELING & DATA VISUALIZATION\n")
    r_title.font.bold = True
    r_title.font.size = Pt(13)
    r_title.font.color.rgb = RGBColor(192, 57, 43)

    # Info Table
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows_data = [
        ("Student Name:", "[Insert Student Name]"),
        ("Roll Number:", "[Insert Roll Number]"),
        ("Date of Experiment:", "[Insert Date]"),
    ]
    for i, (k, v) in enumerate(rows_data):
        r = info_table.rows[i]
        r.cells[0].text = k
        r.cells[0].paragraphs[0].runs[0].font.bold = True
        r.cells[1].text = v
        r.cells[0].width = Inches(2.0)
        r.cells[1].width = Inches(4.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Objectives
    h1 = doc.add_heading("1. Learning Objectives", level=1)
    p1 = doc.add_paragraph(style="List Bullet")
    p1.add_run("Understand data framing and statistical aggregation using Python and Pandas.")
    p2 = doc.add_paragraph(style="List Bullet")
    p2.add_run("Construct visual representations and bivariate scatter trend analyses with Matplotlib.")
    p3 = doc.add_paragraph(style="List Bullet")
    p3.add_run("Implement efficient search algorithms to locate items in sorted datasets.")

    # Lab Tasks
    doc.add_heading("2. Lab Tasks", level=1)

    # Task 1
    doc.add_heading("Task 1: Tabular Data Creation & Summary Statistics", level=2)
    p_t1 = doc.add_paragraph(
        "Create a synthetic student dataset containing StudentID, Quiz_Score, Assignment_Score, and Study_Hours. "
        "Calculate a weighted Total_Score (40% Quiz, 60% Assignment) and print the descriptive statistics table "
        "(mean, std, min, max). Finally, identify the student with the highest total score.\n"
        "Instructions: Write a Python program and provide a terminal execution screenshot."
    )
    doc.add_paragraph("Code / Solution Area:").runs[0].font.italic = True
    doc.add_paragraph("[Insert Code Here]")
    doc.add_paragraph("Screenshot of Output:").runs[0].font.italic = True
    doc.add_paragraph("[Paste Terminal Screenshot Here]")

    # Task 2
    doc.add_heading("Task 2: Study Hours vs Score Correlation Plot", level=2)
    p_t2 = doc.add_paragraph(
        "Generate a scatter plot representing student study hours against exam performance. "
        "Fit a linear regression trend line to model the correlation and save the graphical output to 'plot_output.png'.\n"
        "Discussion Question: What does the slope of the trend line indicate regarding study commitment?"
    )
    doc.add_paragraph("Code / Solution Area:").runs[0].font.italic = True
    doc.add_paragraph("[Insert Code Here]")
    doc.add_paragraph("Screenshot of Output:").runs[0].font.italic = True
    doc.add_paragraph("[Paste Output / Plot Screenshot Here]")

    # Task 3
    doc.add_heading("Task 3: Binary Search for Student Records", level=2)
    p_t3 = doc.add_paragraph(
        "Implement the Binary Search algorithm in Python to search for a specific target value in a sorted list of scores. "
        "Demonstrate the search process and print whether the target was located and at what index.\n"
        "Discussion Question: Why does binary search have logarithmic time complexity O(log n)?"
    )
    doc.add_paragraph("Code / Solution Area:").runs[0].font.italic = True
    doc.add_paragraph("[Insert Code Here]")
    doc.add_paragraph("Screenshot of Output:").runs[0].font.italic = True
    doc.add_paragraph("[Paste Terminal Screenshot Here]")

    doc.save(dest_path)
    print(f"Sample manual generated successfully at: {dest_path}")


if __name__ == "__main__":
    sample_file = Path("/mnt/SSD_Storage/Coding/Lab-Gen/examples/sample_lab_manual.docx")
    create_sample_manual(sample_file)
